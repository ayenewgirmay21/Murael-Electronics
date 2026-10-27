import csv
import io
import logging
import os
import sqlite3
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation

from dotenv import load_dotenv
from telegram import (
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    InputFile,
    Update,
)
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    ConversationHandler,
    MessageHandler,
    filters,
)

load_dotenv()
TOKEN = os.getenv("TELEGRAM_BOT_TOKEN", "").strip()
BUSINESS_NAME = os.getenv("BUSINESS_NAME", "Murael Electronics").strip()
BUSINESS_CONTACT = os.getenv("BUSINESS_CONTACT", "@Murael21bot").strip()
try:
    ADMIN_USER_IDS = {
        int(item.strip())
        for item in os.getenv("ADMIN_USER_IDS", "").split(",")
        if item.strip()
    }
except ValueError:
    raise RuntimeError("ADMIN_USER_IDS must contain numeric Telegram user IDs separated by commas.")

DB_PATH = os.getenv("DATABASE_PATH", "data/murael.db")
CURRENCY = "ETB"

logging.basicConfig(
    format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
    level=logging.INFO,
)
logging.getLogger("httpx").setLevel(logging.WARNING)
logger = logging.getLogger("murael_bot")

# Conversation states for secure admin product creation.
P_NAME, P_CATEGORY, P_PRICE, P_STOCK, P_DESCRIPTION, P_PHOTO = range(6)
P_IMPORT = 6


def db_connect():
    os.makedirs(os.path.dirname(DB_PATH) or ".", exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    with db_connect() as conn:
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS products (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name_en TEXT NOT NULL,
            name_am TEXT NOT NULL DEFAULT '',
            category TEXT NOT NULL,
            price REAL NOT NULL CHECK(price >= 0),
            stock INTEGER NOT NULL DEFAULT 0 CHECK(stock >= 0),
            description_en TEXT NOT NULL DEFAULT '',
            description_am TEXT NOT NULL DEFAULT '',
            photo_file_id TEXT NOT NULL DEFAULT '',
            active INTEGER NOT NULL DEFAULT 1,
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS cart (
            user_id INTEGER NOT NULL,
            product_id INTEGER NOT NULL REFERENCES products(id),
            quantity INTEGER NOT NULL CHECK(quantity > 0),
            PRIMARY KEY(user_id, product_id)
        );
        CREATE TABLE IF NOT EXISTS orders (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            customer_name TEXT NOT NULL,
            phone TEXT NOT NULL,
            language TEXT NOT NULL DEFAULT 'en',
            fulfillment TEXT NOT NULL,
            address TEXT NOT NULL DEFAULT '',
            notes TEXT NOT NULL DEFAULT '',
            total REAL NOT NULL,
            status TEXT NOT NULL DEFAULT 'NEW',
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS order_items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            order_id INTEGER NOT NULL REFERENCES orders(id) ON DELETE CASCADE,
            product_id INTEGER NOT NULL,
            product_name TEXT NOT NULL,
            unit_price REAL NOT NULL,
            quantity INTEGER NOT NULL
        );
        """)
        conn.commit()


TEXT = {
    "en": {
        "welcome": "Welcome to {name}!\nBrowse products, add items to your cart, and send an order request.",
        "menu": "Choose an option:",
        "products": "🛍 Browse products",
        "cart": "🛒 My cart",
        "contact": "☎️ Contact us",
        "language": "🌐 Language",
        "about": "ℹ️ About us",
        "back": "⬅️ Main menu",
        "categories": "Choose a category:",
        "empty": "There are no active products in this category yet. Please contact us.",
        "added": "Added {name} to your cart.",
        "out": "Sorry, this item is currently unavailable or stock is insufficient.",
        "cart_empty": "Your cart is empty.",
        "cart_title": "🛒 Your cart",
        "checkout": "✅ Place order request",
        "clear": "🗑 Clear cart",
        "choose_language": "Choose your language:",
        "lang_saved": "Language updated.",
        "contact_text": "Contact {name}: {contact}\nAsk us about products, prices, stock, or repairs.",
        "support_title": "Customer Support",
        "support_ai": "Ask AI Assistant",
        "support_orders": "Order Assistance",
        "support_delivery": "Delivery Information",
        "support_contact": "Contact the Shop",
        "support_ai_prompt": "Type your question in English or Amharic. The local AI assistant can help with catalogue questions.",
        "support_delivery_text": "Delivery availability, fees, and timing are not confirmed in this catalogue. Please contact the shop before placing an order.",
        "about_text": "{name} sells electronics and accessories and provides repair services. Product availability and final order details are confirmed by our team.",
        "ask_name": "Please type your full name:",
        "ask_phone": "Please type a phone number where we can reach you:",
        "fulfillment": "Choose collection or delivery:",
        "pickup": "🏬 Shop collection",
        "delivery": "🚚 Delivery",
        "ask_address": "Type your delivery address and area:",
        "ask_notes": "Any extra notes? Type your note, or send - for none.",
        "confirm": "Review your request:\n\nName: {name}\nPhone: {phone}\nFulfillment: {fulfillment}{address}\nTotal: {total} ETB\n\nSend this order request?",
        "confirm_yes": "📨 Submit order",
        "confirm_no": "Cancel",
        "order_received": "Thank you! Your order request #{order_id} was sent to {name}. Our team will contact you to confirm availability, delivery, and payment. This is not a payment confirmation.",
        "cancelled": "Order request cancelled. Your cart is unchanged.",
        "order_cancelled": "Order request cancelled.",
        "admin_only": "This command is restricted to authorized administrators.",
        "admin_menu": "🔐 Admin panel\nChoose an action:",
        "add_product": "➕ Add product",
        "orders": "📋 Recent orders",
        "stock": "📦 Update stock",
        "import": "📥 Import CSV/Excel",
        "cancel": "Cancel",
        "ask_product_name": "Enter product name in English:",
        "ask_category": "Enter a category: Phones, Computers, Devices, Components, Repairs, or Other.",
        "ask_price": "Enter price in ETB (example: 1250.00):",
        "ask_stock": "Enter stock quantity as a whole number. For a service, use 999 or the available number of bookings.",
        "ask_description": "Enter the English description (or - for none):",
        "ask_photo": "Send a product photo now, or send - to skip. The photo is stored by Telegram file ID.",
        "product_added": "Product added successfully with ID #{product_id}. English-only fields are supported for now; you can edit the database/CSV to add Amharic translations.",
        "bad_price": "That price is not valid. Enter a number such as 1250 or 1250.50.",
        "bad_stock": "Enter a non-negative whole number, such as 5.",
        "bad_csv": "Could not import this file. Please use the provided CSV template and include required columns: name_en, category, price, stock. Excel .xlsx files are also supported.",
        "imported": "Import complete. Added {count} product(s). Skipped {skipped} row(s).",
        "upload_file": "Upload a .csv or .xlsx file now. Use /cancel to stop.",
        "stock_prompt": "Send stock update in this format: product_id,quantity (example: 3,12).",
        "stock_updated": "Stock updated for product #{product_id}.",
        "stock_error": "Could not update stock. Check the product ID and quantity.",
        "orders_empty": "No orders have been placed yet.",
        "orders_title": "Recent orders:",
        "order_status_prompt": "To update an order, send: /status ORDER_ID STATUS\nAllowed statuses: NEW, CONFIRMED, READY, COMPLETED, CANCELLED",
        "status_updated": "Order #{order_id} status changed to {status}.",
        "help": "Commands:\n/start - Main menu\n/help - Help\n/admin - Admin panel (authorized admins only)\n/cancel - Cancel current admin/order step",
        "unknown": "I didn't understand that. Use /start to open the menu.",
    },
    "am": {
        "welcome": "ወደ {name} እንኳን በደህና መጡ!\nምርቶችን ይመልከቱ፣ ወደ ጋሪ ይጨምሩ፣ የትዕዛዝ ጥያቄ ይላኩ።",
        "menu": "ከታች ያለውን ይምረጡ፦",
        "products": "🛍 ምርቶችን ይመልከቱ",
        "cart": "🛒 ጋሪዬ",
        "contact": "☎️ ያግኙን",
        "language": "🌐 ቋንቋ",
        "about": "ℹ️ ስለ እኛ",
        "back": "⬅️ ዋና ምናሌ",
        "categories": "ምድብ ይምረጡ፦",
        "empty": "በዚህ ምድብ ውስጥ አሁን የሚታይ ምርት የለም። እባክዎ ያግኙን።",
        "added": "{name} ወደ ጋሪዎ ተጨምሯል።",
        "out": "ይቅርታ፣ ይህ ምርት አሁን የለም ወይም በቂ ክምችት የለም።",
        "cart_empty": "ጋሪዎ ባዶ ነው።",
        "cart_title": "🛒 ጋሪዎ",
        "checkout": "✅ የትዕዛዝ ጥያቄ ይላኩ",
        "clear": "🗑 ጋሪውን አጽዳ",
        "choose_language": "ቋንቋዎን ይምረጡ፦",
        "lang_saved": "ቋንቋው ተቀይሯል።",
        "contact_text": "{name} ያግኙ፦ {contact}\nስለ ምርቶች፣ ዋጋ፣ ክምችት ወይም ጥገና ይጠይቁ።",
        "support_title": "የደንበኛ ድጋፍ",
        "support_ai": "የAI እገዛ",
        "support_orders": "ስለ ትዕዛዝ እገዛ",
        "support_delivery": "የማድረስ መረጃ",
        "support_contact": "ሱቁን ያነጋግሩ",
        "support_ai_prompt": "ጥያቄዎን በአማርኛ ወይም በእንግሊዝኛ ይጻፉ። የአካባቢው AI ረዳት ስለ ምርቶች መረጃ ሊረዳዎት ይችላል።",
        "support_delivery_text": "የማድረስ አገልግሎት፣ ክፍያ እና ጊዜ በዚህ መረጃ አልተረጋገጡም። ትዕዛዝ ከማዘዝዎ በፊት ሱቁን ያነጋግሩ።",
        "about_text": "{name} የኤሌክትሮኒክስ እቃዎችና መለዋወጫዎችን ይሸጣል፣ የጥገና አገልግሎትም ይሰጣል። የምርት መኖርና የመጨረሻ የትዕዛዝ ዝርዝር በቡድናችን ይረጋገጣል።",
        "ask_name": "ሙሉ ስምዎን ይጻፉ፦",
        "ask_phone": "ሊደርስብዎ የሚችል ስልክ ቁጥር ይጻፉ፦",
        "fulfillment": "መውሰጃ ወይም ማድረሻ ይምረጡ፦",
        "pickup": "🏬 ከሱቅ መውሰድ",
        "delivery": "🚚 ማድረስ",
        "ask_address": "የማድረሻ አድራሻዎንና አካባቢዎን ይጻፉ፦",
        "ask_notes": "ተጨማሪ ማስታወሻ አለ? ካልሆነ - ይላኩ።",
        "confirm": "የትዕዛዝ ጥያቄዎን ይመልከቱ፦\n\nስም፦ {name}\nስልክ፦ {phone}\nአገልግሎት፦ {fulfillment}{address}\nጠቅላላ፦ {total} ብር\n\nይህን ጥያቄ ይላኩ?",
        "confirm_yes": "📨 ትዕዛዝ ላክ",
        "confirm_no": "ሰርዝ",
        "order_received": "እናመሰግናለን! የትዕዛዝ ጥያቄዎ # {order_id} ተልኳል። ቡድናችን የምርት መኖር፣ ማድረስና ክፍያን ለማረጋገጥ ያነጋግርዎታል። ይህ የክፍያ ማረጋገጫ አይደለም።",
        "cancelled": "የትዕዛዝ ጥያቄው ተሰርዟል። ጋሪዎ አልተቀየረም።",
        "order_cancelled": "የትዕዛዝ ጥያቄው ተሰርዟል።",
        "admin_only": "ይህ ትዕዛዝ ለተፈቀደላቸው አስተዳዳሪዎች ብቻ ነው።",
        "admin_menu": "🔐 የአስተዳዳሪ ምናሌ\nእርምጃ ይምረጡ፦",
        "add_product": "➕ ምርት ጨምር",
        "orders": "📋 የቅርብ ትዕዛዞች",
        "stock": "📦 ክምችት አዘምን",
        "import": "📥 CSV/Excel አስገባ",
        "cancel": "ሰርዝ",
        "ask_product_name": "የምርቱን ስም በእንግሊዝኛ ያስገቡ፦",
        "ask_category": "ምድብ ያስገቡ፦ Phones, Computers, Devices, Components, Repairs ወይም Other።",
        "ask_price": "ዋጋውን በብር ያስገቡ (ምሳሌ፦ 1250.00)፦",
        "ask_stock": "የክምችት ብዛት በሙሉ ቁጥር ያስገቡ። ለጥገና አገልግሎት የሚገኙ ቦታዎችን ያስገቡ።",
        "ask_description": "የምርቱን መግለጫ በእንግሊዝኛ ያስገቡ (ወይም - ለማለፍ)፦",
        "ask_photo": "የምርቱን ፎቶ ይላኩ፣ ወይም - በማለት ይዝለሉ።",
        "product_added": "ምርቱ በተሳካ ሁኔታ ተጨምሯል። መለያ #{product_id}።",
        "bad_price": "ትክክለኛ ዋጋ ያስገቡ፣ ለምሳሌ 1250 ወይም 1250.50።",
        "bad_stock": "ከዜሮ በላይ ወይም ዜሮ የሆነ ሙሉ ቁጥር ያስገቡ።",
        "bad_csv": "ፋይሉን ማስገባት አልተቻለም። የምሳሌ CSV ፋይሉን ይጠቀሙ።",
        "imported": "ማስገባት ተጠናቋል። {count} ምርት(ዎች) ተጨምረዋል፤ {skipped} ረድፎች ተዘለዋል።",
        "upload_file": "አሁን .csv ወይም .xlsx ፋይል ይላኩ። ለማቆም /cancel ይጠቀሙ።",
        "stock_prompt": "በዚህ ቅርጽ ይላኩ፦ product_id,quantity (ምሳሌ፦ 3,12)።",
        "stock_updated": "የምርት #{product_id} ክምችት ተዘምኗል።",
        "stock_error": "ክምችቱን ማዘመን አልተቻለም። የምርት መለያና ብዛት ይፈትሹ።",
        "orders_empty": "እስካሁን ትዕዛዝ አልተላከም።",
        "orders_title": "የቅርብ ትዕዛዞች፦",
        "order_status_prompt": "የትዕዛዝ ሁኔታ ለመቀየር፦ /status ORDER_ID STATUS\nሁኔታዎች፦ NEW, CONFIRMED, READY, COMPLETED, CANCELLED",
        "status_updated": "የትዕዛዝ #{order_id} ሁኔታ ወደ {status} ተቀይሯል።",
        "help": "ትዕዛዞች፦\n/start - ዋና ምናሌ\n/help - እገዛ\n/admin - የአስተዳዳሪ ምናሌ\n/cancel - የአሁኑን ሂደት ሰርዝ",
        "unknown": "/start በመጫን ዋና ምናሌውን ይክፈቱ።",
    },
}


def get_lang(user_id):
    with db_connect() as conn:
        row = conn.execute("SELECT language FROM orders WHERE user_id=? ORDER BY id DESC LIMIT 1", (user_id,)).fetchone()
    # Language preference is stored in Telegram's user_data, not inferred from orders.
    return None


def lang(context):
    return context.user_data.get("language", "en")


def tr(context, key, **kwargs):
    language = lang(context)
    template = TEXT.get(language, TEXT["en"]).get(key, TEXT["en"].get(key, key))
    return template.format(**kwargs)


def is_admin(update: Update):
    return bool(update.effective_user and update.effective_user.id in ADMIN_USER_IDS)


def main_keyboard(context):
    return InlineKeyboardMarkup([
        [InlineKeyboardButton(tr(context, "products"), callback_data="products")],
        [
            InlineKeyboardButton(tr(context, "cart"), callback_data="cart"),
            InlineKeyboardButton(tr(context, "contact"), callback_data="contact"),
        ],
        [InlineKeyboardButton(tr(context, "support_title"), callback_data="support")],
        [
            InlineKeyboardButton(tr(context, "about"), callback_data="about"),
            InlineKeyboardButton(tr(context, "language"), callback_data="language"),
        ],
    ])


def support_keyboard(context):
    return InlineKeyboardMarkup([
        [InlineKeyboardButton(tr(context, "support_ai"), callback_data="support:ai")],
        [InlineKeyboardButton(tr(context, "support_orders"), callback_data="support:orders")],
        [InlineKeyboardButton(tr(context, "support_delivery"), callback_data="support:delivery")],
        [InlineKeyboardButton(tr(context, "support_contact"), callback_data="support:contact")],
        [InlineKeyboardButton(tr(context, "back"), callback_data="home")],
    ])


def categories_keyboard(context):
    names = [
        ("Phones", "📱 Phones / ስልኮች"),
        ("Computers", "💻 Computers / ኮምፒውተሮች"),
        ("Devices", "🔌 Devices / መሳሪያዎች"),
        ("Components", "⚙️ Components / ክፍሎች"),
        ("Repairs", "🛠 Repairs / ጥገና"),
        ("Other", "📦 Other / ሌሎች"),
    ]
    keyboard = [[InlineKeyboardButton(label, callback_data=f"cat:{code}")] for code, label in names]
    keyboard.append([InlineKeyboardButton(tr(context, "back"), callback_data="home")])
    return InlineKeyboardMarkup(keyboard)


def money(value):
    return f"{value:,.2f}"


def safe_text(value):
    return (value or "").strip()


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.effective_message.reply_text(
        tr(context, "welcome", name=BUSINESS_NAME),
        reply_markup=main_keyboard(context),
    )


async def myid_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Show the caller's numeric Telegram user ID for local admin configuration."""
    user = update.effective_user
    if user:
        await update.effective_message.reply_text(
            f"Your Telegram user ID is: {user.id}\n"
            "Add this number to ADMIN_USER_IDS in your local .env file, then restart the bot."
        )


async def help_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.effective_message.reply_text(tr(context, "help"))


async def show_products(query, context, category=None):
    with db_connect() as conn:
        if category:
            rows = conn.execute(
                "SELECT * FROM products WHERE active=1 AND category=? ORDER BY id DESC LIMIT 30",
                (category,),
            ).fetchall()
        else:
            rows = conn.execute("SELECT * FROM products WHERE active=1 ORDER BY id DESC LIMIT 30").fetchall()

    if not rows:
        await query.edit_message_text(
            tr(context, "empty"),
            reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton(tr(context, "back"), callback_data="products")]]),
        )
        return

    # Keep the category menu unchanged; send each product as a separate message.
    for p in rows:
        name = p["name_am"] if lang(context) == "am" and p["name_am"] else p["name_en"]
        desc = p["description_am"] if lang(context) == "am" and p["description_am"] else p["description_en"]
        stock_text = f"Stock: {p['stock']}" if lang(context) == "en" else f"ክምችት፦ {p['stock']}"
        text = f"{name}\n{desc}\nPrice: {money(p['price'])} ETB\n{stock_text}\nID: {p['id']}"
        if lang(context) == "am":
            text = f"{name}\n{desc}\nዋጋ፦ {money(p['price'])} ብር\n{stock_text}\nመለያ፦ {p['id']}"
        markup = InlineKeyboardMarkup([[InlineKeyboardButton("➕ Add / ጨምር", callback_data=f"add:{p['id']}")]])
        try:
            if p["photo_file_id"]:
                await query.message.reply_photo(photo=p["photo_file_id"], caption=text, reply_markup=markup)
            else:
                await query.message.reply_text(text, reply_markup=markup)
        except Exception:
            logger.exception("Could not send product %s", p["id"])
            await query.message.reply_text(text, reply_markup=markup)


def cart_contents(user_id):
    with db_connect() as conn:
        return conn.execute("""
            SELECT c.product_id, c.quantity, p.name_en, p.name_am, p.price, p.stock
            FROM cart c JOIN products p ON p.id=c.product_id
            WHERE c.user_id=? AND p.active=1 ORDER BY p.name_en
        """, (user_id,)).fetchall()


async def show_cart(query, context):
    rows = cart_contents(query.from_user.id)
    if not rows:
        await query.edit_message_text(tr(context, "cart_empty"), reply_markup=InlineKeyboardMarkup([
            [InlineKeyboardButton(tr(context, "products"), callback_data="products")],
            [InlineKeyboardButton(tr(context, "back"), callback_data="home")],
        ]))
        return
    total = sum(float(r["price"]) * int(r["quantity"]) for r in rows)
    lines = [tr(context, "cart_title"), ""]
    for r in rows:
        name = r["name_am"] if lang(context) == "am" and r["name_am"] else r["name_en"]
        lines.append(f"{name} × {r['quantity']} — {money(r['price'] * r['quantity'])} ETB")
    lines += ["", f"Total: {money(total)} ETB"]
    await query.edit_message_text("\n".join(lines), reply_markup=InlineKeyboardMarkup([
        [InlineKeyboardButton(tr(context, "checkout"), callback_data="checkout")],
        [InlineKeyboardButton(tr(context, "clear"), callback_data="clear_cart")],
        [InlineKeyboardButton(tr(context, "products"), callback_data="products")],
        [InlineKeyboardButton(tr(context, "back"), callback_data="home")],
    ]))


async def admin_panel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update):
        await update.effective_message.reply_text(tr(context, "admin_only"))
        return
    await update.effective_message.reply_text(
        tr(context, "admin_menu"),
        reply_markup=InlineKeyboardMarkup([
            [InlineKeyboardButton(tr(context, "add_product"), callback_data="admin:add")],
            [InlineKeyboardButton(tr(context, "orders"), callback_data="admin:orders")],
            [InlineKeyboardButton(tr(context, "stock"), callback_data="admin:stock")],
            [InlineKeyboardButton(tr(context, "import"), callback_data="admin:import")],
        ]),
    )


async def admin_button(query, context):
    if query.from_user.id not in ADMIN_USER_IDS:
        await query.answer("Admin only.", show_alert=True)
        return
    if query.data == "admin:add":
        context.user_data["new_product"] = {}
        await query.message.reply_text(tr(context, "ask_product_name"))
        return P_NAME
    if query.data == "admin:orders":
        with db_connect() as conn:
            orders = conn.execute("SELECT * FROM orders ORDER BY id DESC LIMIT 10").fetchall()
            items_by_order = {}
            for o in orders:
                items_by_order[o["id"]] = conn.execute("SELECT product_name, quantity FROM order_items WHERE order_id=?", (o["id"],)).fetchall()
        if not orders:
            await query.message.reply_text(tr(context, "orders_empty"))
        else:
            await query.message.reply_text(tr(context, "orders_title"))
            for o in orders:
                item_line = ", ".join(f"{i['product_name']} x{i['quantity']}" for i in items_by_order[o["id"]])
                await query.message.reply_text(
                    f"Order #{o['id']} | {o['status']}\n{o['customer_name']} | {o['phone']}\n{o['fulfillment']} {o['address']}\n{item_line}\nTotal: {money(o['total'])} ETB\nUser ID: {o['user_id']}"
                )
        await query.message.reply_text(tr(context, "order_status_prompt"))
        return ConversationHandler.END
    if query.data == "admin:stock":
        await query.message.reply_text(tr(context, "stock_prompt"))
        return P_IMPORT
    if query.data == "admin:import":
        await query.message.reply_text(tr(context, "upload_file"))
        return P_IMPORT
    return ConversationHandler.END


async def add_name(update, context):
    context.user_data["new_product"]["name_en"] = safe_text(update.message.text)
    await update.message.reply_text(tr(context, "ask_category"))
    return P_CATEGORY


async def add_category(update, context):
    valid = {"phones", "computers", "devices", "components", "repairs", "other"}
    value = safe_text(update.message.text).lower()
    if value not in valid:
        await update.message.reply_text("Please enter one of: Phones, Computers, Devices, Components, Repairs, Other.")
        return P_CATEGORY
    context.user_data["new_product"]["category"] = value.title()
    await update.message.reply_text(tr(context, "ask_price"))
    return P_PRICE


async def add_price(update, context):
    try:
        price = Decimal(safe_text(update.message.text))
        if price < 0 or not price.is_finite():
            raise InvalidOperation
    except (InvalidOperation, ValueError):
        await update.message.reply_text(tr(context, "bad_price"))
        return P_PRICE
    context.user_data["new_product"]["price"] = float(price)
    await update.message.reply_text(tr(context, "ask_stock"))
    return P_STOCK


async def add_stock(update, context):
    try:
        stock = int(safe_text(update.message.text))
        if stock < 0:
            raise ValueError
    except ValueError:
        await update.message.reply_text(tr(context, "bad_stock"))
        return P_STOCK
    context.user_data["new_product"]["stock"] = stock
    await update.message.reply_text(tr(context, "ask_description"))
    return P_DESCRIPTION


async def add_description(update, context):
    description = safe_text(update.message.text)
    context.user_data["new_product"]["description_en"] = "" if description == "-" else description
    await update.message.reply_text(tr(context, "ask_photo"))
    return P_PHOTO


async def add_photo(update, context):
    product = context.user_data.get("new_product", {})
    photo_file_id = ""
    if update.message.photo:
        photo_file_id = update.message.photo[-1].file_id
    elif safe_text(update.message.text) != "-":
        await update.message.reply_text("Please send a photo or type - to skip.")
        return P_PHOTO
    product["photo_file_id"] = photo_file_id
    product.setdefault("name_am", "")
    product.setdefault("description_am", "")
    try:
        with db_connect() as conn:
            cur = conn.execute("""
                INSERT INTO products
                (name_en, name_am, category, price, stock, description_en, description_am, photo_file_id, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                product["name_en"], product["name_am"], product["category"], product["price"],
                product["stock"], product["description_en"], product["description_am"],
                product["photo_file_id"], datetime.now(timezone.utc).isoformat()
            ))
            product_id = cur.lastrowid
            conn.commit()
        await update.message.reply_text(tr(context, "product_added", product_id=product_id))
    except Exception:
        logger.exception("Failed to create product")
        await update.message.reply_text("Could not save product. Check the logs and try again.")
    context.user_data.pop("new_product", None)
    return ConversationHandler.END


async def import_file(update, context):
    # The same state handles stock updates and file imports. Text with a comma is a stock update.
    if update.message.text and "," in update.message.text and is_admin(update):
        try:
            product_id_text, stock_text = update.message.text.split(",", 1)
            product_id, stock = int(product_id_text.strip()), int(stock_text.strip())
            if stock < 0:
                raise ValueError
            with db_connect() as conn:
                cur = conn.execute("UPDATE products SET stock=? WHERE id=? AND active=1", (stock, product_id))
                conn.commit()
                if cur.rowcount != 1:
                    raise ValueError
            await update.message.reply_text(tr(context, "stock_updated", product_id=product_id))
        except Exception:
            await update.message.reply_text(tr(context, "stock_error"))
        return ConversationHandler.END

    if not update.message.document:
        await update.message.reply_text(tr(context, "upload_file"))
        return P_IMPORT
    doc = update.message.document
    filename = (doc.file_name or "").lower()
    if not (filename.endswith(".csv") or filename.endswith(".xlsx")):
        await update.message.reply_text(tr(context, "bad_csv"))
        return P_IMPORT
    try:
        tg_file = await context.bot.get_file(doc.file_id)
        bio = io.BytesIO()
        await tg_file.download_to_memory(out=bio)
        raw = bio.getvalue()
        if filename.endswith(".csv"):
            text_data = raw.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text_data))
            rows = list(reader)
        else:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active
            values = list(ws.iter_rows(values_only=True))
            if not values:
                raise ValueError("Empty spreadsheet")
            headers = [str(v).strip() if v is not None else "" for v in values[0]]
            rows = [dict(zip(headers, row)) for row in values[1:]]
        required = {"name_en", "category", "price", "stock"}
        if not rows or not required.issubset(set(rows[0].keys())):
            raise ValueError("Missing required columns")
        count, skipped = 0, 0
        for row in rows:
            try:
                name = safe_text(str(row.get("name_en", "")))
                category = safe_text(str(row.get("category", "Other"))).title()
                price = float(row.get("price"))
                stock = int(float(row.get("stock")))
                if not name or price < 0 or stock < 0:
                    raise ValueError
                with db_connect() as conn:
                    conn.execute("""
                        INSERT INTO products
                        (name_en, name_am, category, price, stock, description_en, description_am, photo_file_id, created_at)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (
                        name, safe_text(str(row.get("name_am", ""))), category, price, stock,
                        safe_text(str(row.get("description_en", ""))), safe_text(str(row.get("description_am", ""))),
                        safe_text(str(row.get("photo_file_id", ""))), datetime.now(timezone.utc).isoformat()
                    ))
                    conn.commit()
                count += 1
            except Exception:
                skipped += 1
        await update.message.reply_text(tr(context, "imported", count=count, skipped=skipped))
    except Exception:
        logger.exception("Product import failed")
        await update.message.reply_text(tr(context, "bad_csv"))
        return P_IMPORT
    return ConversationHandler.END


async def checkout_start(query, context):
    rows = cart_contents(query.from_user.id)
    if not rows:
        await query.edit_message_text(tr(context, "cart_empty"))
        return ConversationHandler.END
    # Validate current stock before asking for details.
    if any(int(r["quantity"]) > int(r["stock"]) for r in rows):
        await query.edit_message_text(tr(context, "out"))
        return ConversationHandler.END
    context.user_data["checkout"] = {}
    await query.message.reply_text(tr(context, "ask_name"))
    return 10


async def checkout_name(update, context):
    name = safe_text(update.message.text)
    if len(name) < 2:
        await update.message.reply_text(tr(context, "ask_name"))
        return 10
    context.user_data["checkout"]["customer_name"] = name
    await update.message.reply_text(tr(context, "ask_phone"))
    return 11


async def checkout_phone(update, context):
    phone = safe_text(update.message.text)
    if len(phone) < 5:
        await update.message.reply_text(tr(context, "ask_phone"))
        return 11
    context.user_data["checkout"]["phone"] = phone
    await update.message.reply_text(tr(context, "fulfillment"), reply_markup=InlineKeyboardMarkup([
        [InlineKeyboardButton(tr(context, "pickup"), callback_data="fulfill:pickup")],
        [InlineKeyboardButton(tr(context, "delivery"), callback_data="fulfill:delivery")],
    ]))
    return 12


async def checkout_fulfillment(update, context):
    query = update.callback_query
    await query.answer()
    if query.data == "fulfill:delivery":
        context.user_data["checkout"]["fulfillment"] = "Delivery"
        await query.message.reply_text(tr(context, "ask_address"))
        return 13
    context.user_data["checkout"]["fulfillment"] = "Shop collection"
    context.user_data["checkout"]["address"] = ""
    await query.message.reply_text(tr(context, "ask_notes"))
    return 14


async def checkout_address(update, context):
    context.user_data["checkout"]["address"] = safe_text(update.message.text)
    await update.message.reply_text(tr(context, "ask_notes"))
    return 14


async def checkout_notes(update, context):
    context.user_data["checkout"]["notes"] = "" if safe_text(update.message.text) == "-" else safe_text(update.message.text)
    details = context.user_data["checkout"]
    rows = cart_contents(update.effective_user.id)
    total = sum(float(r["price"]) * int(r["quantity"]) for r in rows)
    address = f"\nAddress: {details.get('address', '')}" if details.get("address") else ""
    fulfillment = details["fulfillment"]
    details["total"] = total
    details["summary"] = tr(
        context, "confirm",
        name=details["customer_name"], phone=details["phone"],
        fulfillment=fulfillment, address=address, total=money(total)
    )
    await update.message.reply_text(
        details["summary"],
        reply_markup=InlineKeyboardMarkup([
            [InlineKeyboardButton(tr(context, "confirm_yes"), callback_data="order:submit")],
            [InlineKeyboardButton(tr(context, "confirm_no"), callback_data="order:cancel")],
        ]),
    )
    return 15


async def checkout_confirm(update, context):
    query = update.callback_query
    await query.answer()
    if query.data == "order:cancel":
        context.user_data.pop("checkout", None)
        await query.edit_message_text(tr(context, "cancelled"))
        return ConversationHandler.END

    user = query.from_user
    details = context.user_data.get("checkout")
    if not details:
        await query.edit_message_text("Checkout session expired. Please open your cart and try again.")
        return ConversationHandler.END

    rows = cart_contents(user.id)
    if not rows or any(int(r["quantity"]) > int(r["stock"]) for r in rows):
        await query.edit_message_text(tr(context, "out"))
        return ConversationHandler.END

    # Revalidate and reserve stock atomically, so two customers cannot order the same final item.
    try:
        with db_connect() as conn:
            total = 0.0
            for row in rows:
                cur = conn.execute(
                    "UPDATE products SET stock=stock-? WHERE id=? AND active=1 AND stock>=?",
                    (int(row["quantity"]), int(row["product_id"]), int(row["quantity"]))
                )
                if cur.rowcount != 1:
                    raise ValueError("Insufficient stock")
                total += float(row["price"]) * int(row["quantity"])
            cur = conn.execute("""
                INSERT INTO orders
                (user_id, customer_name, phone, language, fulfillment, address, notes, total, status, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'NEW', ?)
            """, (
                user.id, details["customer_name"], details["phone"], lang(context),
                details["fulfillment"], details.get("address", ""), details.get("notes", ""),
                total, datetime.now(timezone.utc).isoformat()
            ))
            order_id = cur.lastrowid
            for row in rows:
                conn.execute("""
                    INSERT INTO order_items (order_id, product_id, product_name, unit_price, quantity)
                    VALUES (?, ?, ?, ?, ?)
                """, (order_id, row["product_id"], row["name_en"], row["price"], row["quantity"]))
            conn.execute("DELETE FROM cart WHERE user_id=?", (user.id,))
            conn.commit()
    except Exception:
        logger.exception("Order submission failed")
        await query.edit_message_text(tr(context, "out"))
        return ConversationHandler.END

    await query.edit_message_text(tr(context, "order_received", order_id=order_id, name=BUSINESS_NAME))
    context.user_data.pop("checkout", None)

    if ADMIN_USER_IDS:
        item_lines = "\n".join(f"• {r['name_en']} × {r['quantity']} — {money(r['price'] * r['quantity'])} ETB" for r in rows)
        admin_message = (
            f"🛒 NEW ORDER #{order_id}\n"
            f"Customer: {details['customer_name']}\nPhone: {details['phone']}\n"
            f"Telegram user ID: {user.id}\nFulfillment: {details['fulfillment']}\n"
            f"Address: {details.get('address', '')}\nNotes: {details.get('notes', '')}\n\n"
            f"{item_lines}\nTotal: {money(total)} ETB\nStatus: NEW"
        )
        for admin_id in ADMIN_USER_IDS:
            try:
                await context.bot.send_message(chat_id=admin_id, text=admin_message)
            except Exception:
                logger.exception("Could not notify admin %s", admin_id)
    return ConversationHandler.END


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.pop("new_product", None)
    context.user_data.pop("checkout", None)
    await update.effective_message.reply_text(tr(context, "order_cancelled"))
    return ConversationHandler.END


async def status_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update):
        await update.effective_message.reply_text(tr(context, "admin_only"))
        return
    if len(context.args) != 2:
        await update.effective_message.reply_text(tr(context, "order_status_prompt"))
        return
    try:
        order_id = int(context.args[0])
        status = context.args[1].upper()
        allowed = {"NEW", "CONFIRMED", "READY", "COMPLETED", "CANCELLED"}
        if status not in allowed:
            raise ValueError
        with db_connect() as conn:
            cur = conn.execute("UPDATE orders SET status=? WHERE id=?", (status, order_id))
            conn.commit()
            if cur.rowcount != 1:
                raise ValueError
            order = conn.execute("SELECT user_id FROM orders WHERE id=?", (order_id,)).fetchone()
        await update.effective_message.reply_text(tr(context, "status_updated", order_id=order_id, status=status))
        try:
            await context.bot.send_message(
                chat_id=order["user_id"],
                text=f"Order #{order_id} status: {status}"
            )
        except Exception:
            logger.exception("Could not notify customer about status change")
    except Exception:
        await update.effective_message.reply_text(tr(context, "order_status_prompt"))


async def callback_router(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    data = query.data or ""
    try:
        await query.answer()
    except Exception as exc:
        # Ignore only Telegram's expired/invalid callback-query response.
        from telegram.error import BadRequest
        if isinstance(exc, BadRequest) and (
            "query is too old" in str(exc).lower()
            or "query id is invalid" in str(exc).lower()
        ):
            logger.debug("Ignoring expired callback query: %s", exc)
        else:
            raise

    if data.startswith("admin:"):
        # The admin actions are routed through this handler; ConversationHandler handles replies.
        if not is_admin(update):
            await query.message.reply_text(tr(context, "admin_only"))
            return ConversationHandler.END
        if data == "admin:add":
            context.user_data["new_product"] = {}
            await query.message.reply_text(tr(context, "ask_product_name"))
            return P_NAME
        if data == "admin:orders":
            with db_connect() as conn:
                orders = conn.execute("SELECT * FROM orders ORDER BY id DESC LIMIT 10").fetchall()
                all_items = {o["id"]: conn.execute("SELECT product_name, quantity FROM order_items WHERE order_id=?", (o["id"],)).fetchall() for o in orders}
            if not orders:
                await query.message.reply_text(tr(context, "orders_empty"))
            else:
                await query.message.reply_text(tr(context, "orders_title"))
                for o in orders:
                    item_line = ", ".join(f"{i['product_name']} x{i['quantity']}" for i in all_items[o["id"]])
                    await query.message.reply_text(f"Order #{o['id']} | {o['status']}\n{o['customer_name']} | {o['phone']}\n{o['fulfillment']} {o['address']}\n{item_line}\nTotal: {money(o['total'])} ETB")
            await query.message.reply_text(tr(context, "order_status_prompt"))
            return ConversationHandler.END
        if data == "admin:stock":
            context.user_data["admin_mode"] = "stock"
            await query.message.reply_text(tr(context, "stock_prompt"))
            return P_IMPORT
        if data == "admin:import":
            context.user_data["admin_mode"] = "import"
            await query.message.reply_text(tr(context, "upload_file"))
            return P_IMPORT
        return ConversationHandler.END

    if data == "home":
        await query.edit_message_text(tr(context, "welcome", name=BUSINESS_NAME), reply_markup=main_keyboard(context))
    elif data == "products":
        await query.edit_message_text(tr(context, "categories"), reply_markup=categories_keyboard(context))
    elif data.startswith("cat:"):
        category = data.split(":", 1)[1]
        await show_products(query, context, category)
    elif data.startswith("add:"):
        product_id = int(data.split(":", 1)[1])
        with db_connect() as conn:
            p = conn.execute("SELECT * FROM products WHERE id=? AND active=1", (product_id,)).fetchone()
            if not p or p["stock"] <= 0:
                await query.message.reply_text(tr(context, "out"))
                return
            existing = conn.execute("SELECT quantity FROM cart WHERE user_id=? AND product_id=?", (query.from_user.id, product_id)).fetchone()
            quantity = (existing["quantity"] if existing else 0) + 1
            if quantity > p["stock"]:
                await query.message.reply_text(tr(context, "out"))
                return
            conn.execute("INSERT INTO cart(user_id, product_id, quantity) VALUES (?, ?, ?) ON CONFLICT(user_id, product_id) DO UPDATE SET quantity=excluded.quantity", (query.from_user.id, product_id, quantity))
            conn.commit()
        name = p["name_am"] if lang(context) == "am" and p["name_am"] else p["name_en"]
        await query.message.reply_text(tr(context, "added", name=name))
    elif data == "cart":
        await show_cart(query, context)
    elif data == "clear_cart":
        with db_connect() as conn:
            conn.execute("DELETE FROM cart WHERE user_id=?", (query.from_user.id,))
            conn.commit()
        await query.edit_message_text(tr(context, "cart_empty"), reply_markup=main_keyboard(context))
    elif data == "contact":
        await query.edit_message_text(tr(context, "contact_text", name=BUSINESS_NAME, contact=BUSINESS_CONTACT), reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton(tr(context, "back"), callback_data="home")]]))
    elif data == "support":
        await query.edit_message_text(
            tr(context, "support_title"),
            reply_markup=support_keyboard(context),
        )
    elif data == "support:ai":
        await query.edit_message_text(
            tr(context, "support_ai_prompt"),
            reply_markup=support_keyboard(context),
        )
    elif data == "support:orders":
        # Reuse the existing cart and checkout workflow.
        await show_cart(query, context)
    elif data == "support:delivery":
        await query.edit_message_text(
            tr(context, "support_delivery_text"),
            reply_markup=support_keyboard(context),
        )
    elif data == "support:contact":
        await query.edit_message_text(
            tr(context, "contact_text", name=BUSINESS_NAME, contact=BUSINESS_CONTACT),
            reply_markup=support_keyboard(context),
        )
    elif data == "about":
        await query.edit_message_text(tr(context, "about_text", name=BUSINESS_NAME), reply_markup=InlineKeyboardMarkup([[InlineKeyboardButton(tr(context, "back"), callback_data="home")]]))
    elif data == "language":
        await query.edit_message_text(tr(context, "choose_language"), reply_markup=InlineKeyboardMarkup([
            [InlineKeyboardButton("English", callback_data="lang:en")],
            [InlineKeyboardButton("አማርኛ", callback_data="lang:am")],
        ]))
    elif data.startswith("lang:"):
        context.user_data["language"] = data.split(":", 1)[1]
        await query.edit_message_text(tr(context, "lang_saved"), reply_markup=main_keyboard(context))
    elif data == "checkout":
        if not cart_contents(query.from_user.id):
            await query.edit_message_text(tr(context, "cart_empty"))
            return
        context.user_data["checkout"] = {}
        await query.message.reply_text(tr(context, "ask_name"))
        # Checkout flow is handled by the dedicated conversation handler state.
        context.user_data["checkout_started"] = True
    elif data in {"fulfill:pickup", "fulfill:delivery", "order:submit", "order:cancel"}:
        # These callbacks are consumed by the checkout ConversationHandler.
        return
    else:
        await query.message.reply_text(tr(context, "unknown"))


async def unknown_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    # Never send product-entry or checkout messages to the AI.
    if (
        context.user_data.get("new_product")
        or context.user_data.get("checkout")
        or (is_admin(update) and context.user_data.get("admin_mode"))
    ):
        await update.effective_message.reply_text(tr(context, "unknown"))
        return

    message = update.effective_message
    user = update.effective_user
    if message is None or not message.text or user is None:
        return

    try:
        from ai_sales import answer_customer

        answer = await answer_customer(
            message.text,
            user.id,
            lang(context),
        )
    except Exception:
        logger.exception("AI handler integration failed.")
        answer = tr(context, "unknown")

    await message.reply_text(answer or tr(context, "unknown"))


async def post_init(application: Application):
    init_db()
    await application.bot.set_my_commands([
        ("start", "Open the main menu"),
        ("help", "Get help"),
        ("admin", "Admin panel (authorized admins only)"),
        ("cancel", "Cancel current step"),
    ])
    logger.info("Database ready. Bot is starting.")


def main():
    if not TOKEN or TOKEN == "YOUR_NEW_TOKEN":
        raise RuntimeError("Set TELEGRAM_BOT_TOKEN in your private .env file.")
    if not ADMIN_USER_IDS:
        logger.warning("ADMIN_USER_IDS is empty. Admin features will be unavailable until configured.")
    init_db()
    app = Application.builder().token(TOKEN).post_init(post_init).build()

    # Product creation conversation, initiated by an admin's inline button.
    admin_conv = ConversationHandler(
        entry_points=[CallbackQueryHandler(callback_router, pattern=r"^admin:(add|stock|import|orders)$")],
        states={
            P_NAME: [MessageHandler(filters.TEXT & ~filters.COMMAND, add_name)],
            P_CATEGORY: [MessageHandler(filters.TEXT & ~filters.COMMAND, add_category)],
            P_PRICE: [MessageHandler(filters.TEXT & ~filters.COMMAND, add_price)],
            P_STOCK: [MessageHandler(filters.TEXT & ~filters.COMMAND, add_stock)],
            P_DESCRIPTION: [MessageHandler(filters.TEXT & ~filters.COMMAND, add_description)],
            P_PHOTO: [MessageHandler(filters.PHOTO | (filters.TEXT & ~filters.COMMAND), add_photo)],
            P_IMPORT: [MessageHandler((filters.TEXT & ~filters.COMMAND) | filters.Document.ALL, import_file)],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        allow_reentry=True,
    )

    # Checkout flow. Entry is callback from cart; the first name prompt is issued here.
    checkout_conv = ConversationHandler(
        entry_points=[CallbackQueryHandler(checkout_entry, pattern=r"^checkout$")],
        states={
            10: [MessageHandler(filters.TEXT & ~filters.COMMAND, checkout_name)],
            11: [MessageHandler(filters.TEXT & ~filters.COMMAND, checkout_phone)],
            12: [CallbackQueryHandler(checkout_fulfillment, pattern=r"^fulfill:(pickup|delivery)$")],
            13: [MessageHandler(filters.TEXT & ~filters.COMMAND, checkout_address)],
            14: [MessageHandler(filters.TEXT & ~filters.COMMAND, checkout_notes)],
            15: [CallbackQueryHandler(checkout_confirm, pattern=r"^order:(submit|cancel)$")],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        allow_reentry=True,
    )

    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("help", help_command))
    app.add_handler(CommandHandler("myid", myid_command))
    app.add_handler(CommandHandler("admin", admin_panel))
    app.add_handler(CommandHandler("status", status_command))
    app.add_handler(CommandHandler("cancel", cancel))
    app.add_handler(admin_conv)
    app.add_handler(checkout_conv)
    app.add_handler(CallbackQueryHandler(callback_router))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, unknown_text))
    app.run_polling()


async def checkout_entry(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    rows = cart_contents(query.from_user.id)
    if not rows:
        await query.edit_message_text(tr(context, "cart_empty"))
        return ConversationHandler.END
    if any(int(r["quantity"]) > int(r["stock"]) for r in rows):
        await query.edit_message_text(tr(context, "out"))
        return ConversationHandler.END
    context.user_data["checkout"] = {}
    await query.message.reply_text(tr(context, "ask_name"))
    return 10


if __name__ == "__main__":
    main()
